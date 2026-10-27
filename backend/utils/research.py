"""
Research Test mode (pretest / posttest for a research study).

Flow
----
1. The teacher creates a study (one pretest + one posttest), edits the questions,
   assigns a test to classes (optionally only some students per day) and opens it.
2. A student gives consent, then answers every question. Each answer is ONE short
   recording, uploaded straight to permanent storage (bucket "research-recordings").
   Students never see a score.
3. Background jobs (queues "research_transcribe" and "research_score"):
   * transcribe each recording with Groq Whisper (language auto-detected, so
     Indonesian is kept as Indonesian instead of being translated), and measure
     speaking rate, pauses, response latency and Indonesian words;
   * when all recordings of an attempt are transcribed, Gemini LISTENS to the audio
     and scores Brown's rubric (5 aspects x 1-5) with one fixed model at
     temperature 0. With RESEARCH_SCORING_RUNS > 1 the median of the runs is used.
   Both jobs only use free quota: when it is used up they stop and the sweeper
   (main.py) picks them up again after the reset. Nothing is lost.
4. Human scorers (accounts made by the teacher) listen blind (anonymous codes,
   no pre/post label) and score with the same rubric. AI scores are shown to a
   scorer only after they saved their own.

All database access uses the service role; the API layer checks who is asking.
"""
import base64
import io
import json
import math
import os
import re
import secrets
import statistics
import time
import zipfile
from datetime import datetime, timedelta, timezone
from typing import Dict, List, Optional
from zoneinfo import ZoneInfo

import httpx

from utils import ai

r = ai.redis_conn

BUCKET = os.getenv("RESEARCH_BUCKET", "research-recordings")
GEMINI_MODEL = os.getenv("RESEARCH_GEMINI_MODEL", "gemini-3.8-flash")
GEMINI_RPM = int(os.getenv("RESEARCH_GEMINI_RPM", "5"))
GEMINI_RPD = int(os.getenv("RESEARCH_GEMINI_RPD", "20"))   # replaced by the real value when Google reports it
SCORING_RUNS = max(1, int(os.getenv("RESEARCH_SCORING_RUNS", "1")))
STORAGE_LIMIT_MB = int(os.getenv("RESEARCH_STORAGE_LIMIT_MB", "1024"))
PROMPT_VERSION = "brown-v1"
ASPECTS = ["pronunciation", "fluency", "vocabulary", "grammar", "comprehension"]
ASPECT_LABELS = {
    "pronunciation": "Pronunciation", "fluency": "Fluency", "vocabulary": "Vocabulary",
    "grammar": "Grammar", "comprehension": "Comprehension / Interaction",
}
PACIFIC = ZoneInfo("America/Los_Angeles")   # Gemini daily quotas reset at midnight Pacific
LOCAL_TZ = ZoneInfo(os.getenv("RESEARCH_LOCAL_TZ", "Asia/Makassar"))
TRANSCRIBE_JOB_DEADLINE = 120   # seconds a transcription job waits for Groq before yielding
MAX_SCORE_FAILURES = 4


class ResearchError(Exception):
    def __init__(self, msg, status=400):
        super().__init__(msg)
        self.status = status


def now_iso():
    return datetime.now(timezone.utc).isoformat()


# =============================================================================
# Supabase (REST + Storage) with the service role
# =============================================================================

def _client(timeout=30):
    force_ipv4 = os.getenv("FORCE_IPV4", "true").lower() == "true"
    transport = httpx.HTTPTransport(local_address="0.0.0.0", retries=2) if force_ipv4 else None
    return httpx.Client(timeout=timeout, transport=transport)


def _headers(extra=None):
    key = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
    return {"apikey": key, "Authorization": f"Bearer {key}", **(extra or {})}


def _rest(path):
    return f"{os.getenv('SUPABASE_URL')}/rest/v1/{path}"


def _check(resp, what):
    if resp.status_code >= 300:
        raise RuntimeError(f"{what} failed ({resp.status_code}): {resp.text[:300]}")


def db_select(table, params, client=None) -> List[dict]:
    """Every row of a PostgREST query (pages through big results)."""
    own = client is None
    client = client or _client()
    try:
        rows, offset, page = [], 0, 1000
        while True:
            resp = client.get(_rest(table), headers=_headers(),
                              params={**params, "limit": page, "offset": offset})
            _check(resp, f"Reading {table}")
            batch = resp.json()
            rows.extend(batch)
            if len(batch) < page:
                return rows
            offset += page
    finally:
        if own:
            client.close()


def db_one(table, params, client=None) -> Optional[dict]:
    rows = db_select(table, {**params}, client)
    return rows[0] if rows else None


def db_insert(table, row, *, upsert_on=None, client=None):
    own = client is None
    client = client or _client()
    try:
        prefer = "return=representation"
        params = {}
        if upsert_on:
            prefer += ",resolution=merge-duplicates"
            params["on_conflict"] = upsert_on
        resp = client.post(_rest(table), headers=_headers({"Prefer": prefer, "Content-Type": "application/json"}),
                           params=params, json=row)
        _check(resp, f"Saving to {table}")
        data = resp.json()
        return data[0] if isinstance(data, list) and data and not isinstance(row, list) else data
    finally:
        if own:
            client.close()


def db_update(table, filters, fields, client=None):
    own = client is None
    client = client or _client()
    try:
        resp = client.patch(_rest(table), headers=_headers({"Prefer": "return=representation",
                                                            "Content-Type": "application/json"}),
                            params=filters, json=fields)
        _check(resp, f"Updating {table}")
        return resp.json()
    finally:
        if own:
            client.close()


def db_delete(table, filters, client=None):
    own = client is None
    client = client or _client()
    try:
        resp = client.delete(_rest(table), headers=_headers(), params=filters)
        _check(resp, f"Deleting from {table}")
    finally:
        if own:
            client.close()


def in_list(values):
    vals = [str(v) for v in values if v]
    return f"in.({','.join(vals)})" if vals else "in.(00000000-0000-0000-0000-000000000000)"


def _storage(path=""):
    return f"{os.getenv('SUPABASE_URL')}/storage/v1{path}"


def _ensure_bucket(client):
    resp = client.post(_storage("/bucket"), headers=_headers(),
                       json={"id": BUCKET, "name": BUCKET, "public": False})
    return resp.status_code in (200, 201, 409) or "already exists" in resp.text.lower()


def storage_upload(path, data: bytes, content_type):
    with _client(60) as client:
        url = _storage(f"/object/{BUCKET}/{path}")
        headers = _headers({"Content-Type": content_type or "application/octet-stream", "x-upsert": "true"})
        resp = client.post(url, headers=headers, content=data)
        if resp.status_code in (400, 404) and "bucket" in resp.text.lower() and _ensure_bucket(client):
            resp = client.post(url, headers=headers, content=data)
        _check(resp, "Uploading recording")


def storage_download(path, client=None) -> bytes:
    own = client is None
    client = client or _client(60)
    try:
        resp = client.get(_storage(f"/object/{BUCKET}/{path}"), headers=_headers())
        _check(resp, "Downloading recording")
        return resp.content
    finally:
        if own:
            client.close()


def storage_delete(paths):
    paths = [p for p in paths if p]
    if not paths:
        return
    with _client() as client:
        for i in range(0, len(paths), 100):
            client.request("DELETE", _storage(f"/object/{BUCKET}"), headers=_headers(),
                           json={"prefixes": paths[i:i + 100]})


def signed_urls(paths, expires_in=3 * 3600) -> Dict[str, str]:
    """Temporary playable links for many objects at once."""
    paths = [p for p in dict.fromkeys(paths) if p]
    if not paths:
        return {}
    out = {}
    with _client() as client:
        for i in range(0, len(paths), 100):
            chunk = paths[i:i + 100]
            resp = client.post(_storage(f"/object/sign/{BUCKET}"), headers=_headers(),
                               json={"expiresIn": int(expires_in), "paths": chunk})
            if resp.status_code != 200:
                continue
            for item in resp.json():
                signed = item.get("signedURL") or item.get("signedUrl")
                if signed and item.get("path"):
                    out[item["path"]] = _storage(signed)
    return out


# =============================================================================
# Default test content (from INSTRUMEN PENELITIAN, Unit SMA)
# =============================================================================

DEFAULT_INTRO = ("You will have a short speaking test with three parts. Please answer in English as much as "
                 "you can. It is okay to make mistakes. Just try your best.")


def default_consent(study_title):
    return (
        f"Tes speaking ini adalah bagian dari penelitian \"{study_title}\".\n\n"
        "Sebelum mulai, bacalah dengan teliti:\n"
        "1. Suaramu akan direkam selama tes.\n"
        "2. Rekaman akan ditranskripsi oleh layanan AI Groq dan dinilai oleh AI Google Gemini. "
        "Layanan AI gratis ini dapat menggunakan data yang dikirim untuk meningkatkan produk mereka, "
        "dan data tersebut dapat ditinjau oleh petugas penyedia layanan.\n"
        "3. Rekaman juga dapat didengarkan oleh penilai (guru) yang ditunjuk oleh peneliti.\n"
        "4. Nilai tidak ditampilkan kepadamu. Hasil tes hanya digunakan untuk penelitian ini.\n"
        "5. Rekaman disimpan sampai penelitian selesai, lalu dihapus.\n\n"
        "Jika kamu tidak setuju, pilih \"Saya tidak setuju\" dan beri tahu gurumu."
    )


def _part1(questions):
    return {
        "id": "p1", "title": "Part 1 - Personal Questions",
        "instructions": "Listen to each question and answer it.",
        "items": [{"id": f"p1q{i + 1}", "prompt": q, "speak_seconds": 30, "prep_seconds": 0}
                  for i, q in enumerate(questions)],
    }


def _part2(scene):
    return {
        "id": "p2", "title": "Part 2 - Picture Description",
        "instructions": "Look at the picture carefully, then answer the questions.",
        "image_path": None,
        "image_note": f"Gambar suasana {scene} saat jam istirahat (berwarna; 6-8 orang dengan 4-5 aktivitas berbeda).",
        "items": [
            {"id": "p2q1", "prompt": "Describe what you see in this picture.", "speak_seconds": 45, "prep_seconds": 10},
            {"id": "p2q2", "prompt": "What do you think is happening, and how do the people feel? Why do you think so?",
             "speak_seconds": 45, "prep_seconds": 0},
        ],
    }


def _part3(main_q, follow_q):
    return {
        "id": "p3", "title": "Part 3 - Short Talk",
        "instructions": "You have 30 seconds to think before you answer the main question.",
        "items": [
            {"id": "p3q1", "prompt": main_q, "speak_seconds": 60, "prep_seconds": 30},
            {"id": "p3q2", "prompt": follow_q, "speak_seconds": 30, "prep_seconds": 0},
        ],
    }


def default_content(kind, study_title):
    if kind == "pre":
        parts = [
            _part1(["Please introduce yourself.", "What do you want to be in the future? Why?",
                    "How do you usually spend your free time, and why?",
                    "What is the most interesting thing you learned this year?"]),
            _part2("kantin sekolah"),
            _part3("Should students use AI for homework? Give your reasons and an example.",
                   "What rules should schools make about using AI?"),
        ]
    else:
        parts = [
            _part1(["Please introduce yourself.", "Which university major would you like to choose? Why?",
                    "Describe a person who inspires you.", "What is one skill you want to improve? Why?"]),
            _part2("halaman sekolah"),
            _part3("Should students be allowed to bring phones to class? Give your reasons and an example.",
                   "What rules should schools make about phones?"),
        ]
    return {"version": 1, "consent_text": default_consent(study_title), "intro": DEFAULT_INTRO,
            "max_repeats": 1, "parts": parts}


def all_items(content) -> List[dict]:
    items = []
    for part in (content or {}).get("parts", []):
        for item in part.get("items", []):
            items.append({**item, "part_id": part.get("id"), "part_title": part.get("title"),
                          "image_path": part.get("image_path")})
    return items


_ID_RE = re.compile(r"^[A-Za-z0-9_-]{1,40}$")


def clean_content(content) -> dict:
    """Validate a test edited by the teacher."""
    if not isinstance(content, dict):
        raise ResearchError("Invalid test content.")
    parts = content.get("parts")
    if not isinstance(parts, list) or not parts:
        raise ResearchError("A test needs at least one part.")
    seen = set()
    out_parts = []
    for p_idx, part in enumerate(parts):
        pid = str(part.get("id") or f"p{p_idx + 1}")
        if not _ID_RE.match(pid):
            raise ResearchError(f"Invalid part id: {pid}")
        items = part.get("items") or []
        if not items:
            raise ResearchError(f"{part.get('title') or pid} has no questions.")
        out_items = []
        for i_idx, item in enumerate(items):
            iid = str(item.get("id") or f"{pid}q{i_idx + 1}")
            if not _ID_RE.match(iid) or iid in seen:
                raise ResearchError(f"Question ids must be unique (problem: {iid}).")
            seen.add(iid)
            prompt = str(item.get("prompt") or "").strip()
            if not prompt:
                raise ResearchError(f"A question in {part.get('title') or pid} is empty.")
            speak = int(item.get("speak_seconds") or 30)
            prep = int(item.get("prep_seconds") or 0)
            if not (10 <= speak <= 300) or not (0 <= prep <= 120):
                raise ResearchError("Answer time must be 10-300 s and thinking time 0-120 s.")
            out_items.append({"id": iid, "prompt": prompt[:600], "speak_seconds": speak, "prep_seconds": prep})
        image_path = part.get("image_path")
        if image_path is not None and (not isinstance(image_path, str) or ".." in image_path):
            image_path = None
        out_parts.append({
            "id": pid, "title": str(part.get("title") or f"Part {p_idx + 1}")[:120],
            "instructions": str(part.get("instructions") or "")[:600],
            "image_path": image_path, "image_note": str(part.get("image_note") or "")[:300],
            "items": out_items,
        })
    return {
        "version": 1,
        "consent_text": str(content.get("consent_text") or "").strip()[:4000],
        "intro": str(content.get("intro") or DEFAULT_INTRO).strip()[:1000],
        "max_repeats": max(0, min(3, int(content.get("max_repeats", 1)))),
        "parts": out_parts,
    }


# =============================================================================
# Anonymous codes
# =============================================================================

_CODE_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"


def new_anon_code():
    return "R-" + "".join(secrets.choice(_CODE_ALPHABET) for _ in range(5))


# =============================================================================
# Transcription + audio measures
# =============================================================================

# Common Indonesian words that are not English words (for the "used Indonesian" measure).
_INDONESIAN = set("""
yang dan saya aku kamu itu ini tidak nggak gak enggak ngga ada dengan untuk karena karna apa bisa juga sudah
udah mau kalau kalo jadi dia mereka kita kami dari ke di pada akan atau tapi tetapi sama lagi sangat banyak
orang sekolah teman bagaimana gimana kenapa mengapa siapa dimana kapan iya terus gitu begitu seperti kayak
lebih punya makan belajar rumah waktu suka cita menjadi ingin pengen lah kan sih dong deh bahasa inggris
pelajaran guru murid siswa tugas soal jawab menurut harus boleh jangan semua setiap biasanya kadang saat
sedang lagi juga tentang hal lain sekali banget aja saja hanya cuma masih belum pernah selalu nanti tahun
hari ini besok kemarin karena bisa nya apakah adalah merupakan dalam oleh bahwa agar supaya ketika
""".split())
_WORD_RE = re.compile(r"[A-Za-z']+")


def _words(text):
    return _WORD_RE.findall(text or "")


def _transcribe_research(audio: bytes, filename: str, duration_seconds: Optional[float]):
    """Groq Whisper with timestamps; language auto-detected. Raises ai.AIBusyError when out of quota."""
    seconds = duration_seconds if duration_seconds and duration_seconds > 0 else len(audio) / 4000
    seconds = max(10, int(seconds) + 1)
    info = {}

    def call(model):
        kwargs = dict(file=(filename, audio), model=model, response_format="verbose_json")
        try:
            resp = ai.get_client().audio.transcriptions.create(
                **kwargs, timestamp_granularities=["word", "segment"])
        except TypeError:  # older SDK: segment timestamps only (still enough for pauses)
            resp = ai.get_client().audio.transcriptions.create(**kwargs)
        except Exception as e:  # the API refused word timestamps: use segments only
            if "timestamp" not in str(e).lower():
                raise
            resp = ai.get_client().audio.transcriptions.create(**kwargs)
        info["model"] = model
        if hasattr(resp, "model_dump"):
            data = resp.model_dump()
        elif isinstance(resp, dict):
            data = resp
        else:
            data = {"text": getattr(resp, "text", "")}
        return data, 0

    data = ai.run_with_fallback("transcribe", ai.TRANSCRIBE_MODELS, call,
                                deadline=time.time() + TRANSCRIBE_JOB_DEADLINE, audio_seconds=seconds)
    return data, info.get("model")


def measure(data: dict, duration_seconds: Optional[float]) -> dict:
    """Speaking rate, pauses, latency and Indonesian words from a verbose Whisper result."""
    text = (data.get("text") or "").strip()
    segments = data.get("segments") or []
    words_ts = data.get("words") or []
    tokens = _words(text)
    n_words = len(tokens)
    indo = sum(1 for t in tokens if t.lower() in _INDONESIAN)

    # Speech timeline: words if Groq returned them, otherwise segments.
    spans = [(float(w.get("start", 0)), float(w.get("end", 0))) for w in words_ts if w.get("end") is not None]
    if not spans:
        spans = [(float(s.get("start", 0)), float(s.get("end", 0))) for s in segments if s.get("end") is not None]
    spans.sort()
    pauses = []
    for (s0, e0), (s1, e1) in zip(spans, spans[1:]):
        gap = s1 - e0
        if gap >= 0.5:
            pauses.append(round(gap, 2))
    first_start = spans[0][0] if spans else None
    last_end = spans[-1][1] if spans else None
    total = float(duration_seconds) if duration_seconds else (last_end or 0)
    speaking_span = (last_end - first_start) if spans and last_end > first_start else 0
    logprobs = [s.get("avg_logprob") for s in segments if isinstance(s.get("avg_logprob"), (int, float))]
    return {
        "duration_s": round(total, 1),
        "words": n_words,
        "speech_rate_wpm": round(n_words / (total / 60), 1) if total > 0 else 0,
        "articulation_wpm": round(n_words / (speaking_span / 60), 1) if speaking_span > 0 else 0,
        "pauses_over_1s": sum(1 for p in pauses if p >= 1.0),
        "pauses_over_2s": sum(1 for p in pauses if p >= 2.0),
        "mean_pause_s": round(statistics.mean(pauses), 2) if pauses else 0,
        "longest_pause_s": max(pauses) if pauses else 0,
        "response_latency_s": round(first_start, 2) if first_start is not None else None,
        "indonesian_words": indo,
        "indonesian_ratio": round(indo / n_words, 3) if n_words else 0,
        "asr_confidence": round(math.exp(statistics.mean(logprobs)), 3) if logprobs else None,
        "language": data.get("language"),
        "timing_source": "words" if words_ts else ("segments" if segments else "none"),
    }


def _lock(key, ttl):
    return bool(r.set(f"research:lock:{key}", "1", nx=True, ex=ttl))


def _unlock(key):
    r.delete(f"research:lock:{key}")


def job_transcribe(response_id):
    try:
        row = db_one("research_responses", {"id": f"eq.{response_id}", "select": "*"})
        if not row or row.get("transcribe_status") == "done":
            return
        try:
            audio = storage_download(row["audio_path"])
            ext = os.path.splitext(row["audio_path"])[1] or ".webm"
            data, model = _transcribe_research(audio, f"answer{ext}", row.get("duration_seconds"))
        except ai.BadAudioError as e:
            db_update("research_responses", {"id": f"eq.{response_id}"},
                      {"transcribe_status": "error", "transcribe_error": f"Unusable audio: {e}"[:300],
                       "transcript": "", "next_try_at": None})
            return
        except ai.AIBusyError as e:
            wait = max(60, int(getattr(e, "retry_after", 0) or 0))
            db_update("research_responses", {"id": f"eq.{response_id}"},
                      {"transcribe_error": "Waiting for free transcription quota.",
                       "next_try_at": (datetime.now(timezone.utc) + timedelta(seconds=wait)).isoformat()})
            return
        except Exception as e:  # network / server problem: try again in a few minutes
            db_update("research_responses", {"id": f"eq.{response_id}"},
                      {"transcribe_error": str(e)[:300],
                       "next_try_at": (datetime.now(timezone.utc) + timedelta(minutes=5)).isoformat()})
            return
        m = measure(data, row.get("duration_seconds"))
        db_update("research_responses", {"id": f"eq.{response_id}"}, {
            "transcript": (data.get("text") or "").strip(),
            "transcript_meta": {**m, "model": model},
            "transcribe_status": "done", "transcribe_error": None, "next_try_at": None,
        })
    finally:
        _unlock(f"tr:{response_id}")


# =============================================================================
# Gemini quota bookkeeping
# =============================================================================

def _pacific_day():
    return datetime.now(PACIFIC).strftime("%Y-%m-%d")


def next_gemini_reset() -> datetime:
    now = datetime.now(PACIFIC)
    tomorrow = (now + timedelta(days=1)).replace(hour=0, minute=0, second=5, microsecond=0)
    return tomorrow.astimezone(timezone.utc)


def gemini_rpd():
    learned = r.get(f"research:gemini:{GEMINI_MODEL}:rpd")
    try:
        return int(learned) if learned else GEMINI_RPD
    except ValueError:
        return GEMINI_RPD


def gemini_used_today():
    return int(r.get(f"research:gemini:{GEMINI_MODEL}:day:{_pacific_day()}") or 0)


def _gemini_blocked_until():
    v = r.get(f"research:gemini:{GEMINI_MODEL}:blocked_until")
    return float(v) if v else 0.0


def _gemini_acquire():
    """Returns 0 when a request may be sent now, else seconds to wait (>= 3600 means 'quota for today used')."""
    blocked = _gemini_blocked_until()
    if blocked > time.time():
        return blocked - time.time()
    day_key = f"research:gemini:{GEMINI_MODEL}:day:{_pacific_day()}"
    if int(r.get(day_key) or 0) >= gemini_rpd():
        return (next_gemini_reset() - datetime.now(timezone.utc)).total_seconds()
    minute = int(time.time() // 60)
    min_key = f"research:gemini:{GEMINI_MODEL}:min:{minute}"
    if int(r.get(min_key) or 0) >= GEMINI_RPM:
        return 61 - (time.time() % 60)
    pipe = r.pipeline()
    pipe.incr(day_key)
    pipe.expire(day_key, 3 * 86400)
    pipe.incr(min_key)
    pipe.expire(min_key, 180)
    pipe.execute()
    return 0


class _GeminiQuota(Exception):
    def __init__(self, wait, daily):
        super().__init__("Gemini quota")
        self.wait = wait
        self.daily = daily


def _handle_429(resp):
    """Learn from Google's quota error: which quota, its size, how long to wait."""
    daily, wait, quota_value = False, 60.0, None
    try:
        details = resp.json().get("error", {}).get("details", [])
    except Exception:
        details = []
    for d in details:
        t = d.get("@type", "")
        if t.endswith("QuotaFailure"):
            for v in d.get("violations", []):
                qid = (v.get("quotaId") or "") + (v.get("quotaMetric") or "")
                if "PerDay" in qid or "per_day" in qid.lower():
                    daily = True
                    quota_value = v.get("quotaValue") or quota_value
        if t.endswith("RetryInfo"):
            m = re.match(r"([\d.]+)s", str(d.get("retryDelay", "")))
            if m:
                wait = float(m.group(1)) + 1
    if daily:
        if quota_value and str(quota_value).isdigit():
            r.set(f"research:gemini:{GEMINI_MODEL}:rpd", int(quota_value))
        until = next_gemini_reset().timestamp()
        r.set(f"research:gemini:{GEMINI_MODEL}:blocked_until", until, ex=2 * 86400)
        return _GeminiQuota(until - time.time(), True)
    return _GeminiQuota(wait, False)


def check_gemini_model():
    """Is the configured model reachable with this key? (uses no quota)"""
    key = os.getenv("GEMINI_API_KEY")
    if not key:
        return {"ok": False, "message": "GEMINI_API_KEY is not set on the server."}
    try:
        with _client(15) as client:
            resp = client.get(f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}",
                              headers={"x-goog-api-key": key})
        if resp.status_code == 200:
            return {"ok": True, "message": f"{GEMINI_MODEL} is available.",
                    "display_name": resp.json().get("displayName")}
        return {"ok": False, "message": f"{GEMINI_MODEL}: {resp.status_code} {resp.text[:200]}"}
    except Exception as e:
        return {"ok": False, "message": f"Could not reach Gemini: {e}"}


# =============================================================================
# Gemini scoring (Brown rubric, listening to the audio)
# =============================================================================

RUBRIC = """
Score each aspect from 1 to 5 using EXACTLY these descriptors (Brown, 2004; the school's research rubric).
Descriptors are given in Indonesian (original) with English in brackets.

PRONUNCIATION (Pelafalan)
5 Jelas dan tepat, hampir tanpa kesalahan [clear and accurate, almost no errors]
4 Sedikit salah ucap, tidak mengganggu pemahaman [a few mispronunciations, do not hinder understanding]
3 Ada salah ucap, masih dapat dipahami [some mispronunciations, still understandable]
2 Sering salah ucap, sebagian sulit dipahami [frequent mispronunciations, partly hard to understand]
1 Sangat sulit dipahami [very hard to understand]

FLUENCY (Kelancaran)
5 Lancar dengan jeda alami [fluent with natural pauses]
4 Cukup lancar, sesekali berhenti mencari kata [fairly fluent, occasionally stops to search for words]
3 Sering berhenti sejenak, pesan tetap tersampaikan [often pauses briefly, message still conveyed]
2 Sering berhenti lama, kalimat terputus [often long pauses, broken sentences]
1 Sangat terputus-putus, banyak diam [very fragmented, a lot of silence]

VOCABULARY (Kosakata)
5 Bervariasi dan tepat sesuai topik [varied and appropriate to the topic]
4 Cukup bervariasi, sesekali kurang tepat [fairly varied, occasionally inappropriate]
3 Terbatas tetapi cukup untuk ide sederhana [limited but enough for simple ideas]
2 Sangat terbatas, sering memakai bahasa Indonesia [very limited, often uses Indonesian]
1 Hampir tidak mampu memakai kosakata bahasa Inggris [almost unable to use English vocabulary]

GRAMMAR (Tata bahasa)
5 Hampir tanpa kesalahan [almost no errors]
4 Sedikit kesalahan, tidak mengganggu makna [few errors, meaning not affected]
3 Beberapa kesalahan, makna masih jelas [some errors, meaning still clear]
2 Banyak kesalahan, makna sering kabur [many errors, meaning often unclear]
1 Kesalahan sangat banyak, makna tidak tersampaikan [very many errors, meaning not conveyed]

COMPREHENSION / INTERACTION (Pemahaman dan interaksi)
5 Memahami semua pertanyaan, merespons tepat dan langsung [understands all questions, responds appropriately and promptly]
4 Memahami hampir semua pertanyaan, respons tepat [understands almost all questions, responses appropriate]
3 Kadang perlu pengulangan pertanyaan [sometimes needs the question repeated]
2 Sering perlu pengulangan, respons kurang sesuai [often needs repetition, responses not quite relevant]
1 Tidak memahami pertanyaan [does not understand the questions]
""".strip()

SCORE_SCHEMA = {
    "type": "OBJECT",
    "properties": {
        **{a: {"type": "OBJECT", "properties": {
            "score": {"type": "INTEGER"},
            "evidence": {"type": "STRING"},
        }, "required": ["score", "evidence"]} for a in ASPECTS},
        "overall_comment": {"type": "STRING"},
    },
    "required": ASPECTS + ["overall_comment"],
}

_GEMINI_MIME = {"audio/mp4": "audio/m4a", "audio/x-m4a": "audio/m4a", "audio/mpeg": "audio/mp3"}


def _scoring_parts(content, responses, audio_by_item):
    by_item = {r_["item_id"]: r_ for r_ in responses}
    max_rep = (content or {}).get("max_repeats", 1)
    parts = [{"text": (
        "You are a strict, consistent English speaking examiner for Indonesian senior high school students "
        "(the school is Indonesian; students learn English as a foreign language). This is a research "
        "pretest/posttest, so apply the rubric the same way to every student.\n\n"
        "The student took a recorded speaking test. For each question you get the question, the student's "
        "AUDIO answer, an automatic transcript (it can contain errors - always trust the audio over the "
        "transcript), and measurements taken from the audio. The student could hear each question again "
        f"at most {max_rep} time(s); 'repeats used' tells you how often they asked for it.\n\n"
        f"{RUBRIC}\n\n"
        "Rules:\n"
        "- LISTEN to the audio for Pronunciation and Fluency. Judge pronunciation by intelligibility for an "
        "international listener, not by native-like accent.\n"
        "- Speech in Indonesian does not count as English vocabulary; frequent Indonesian lowers Vocabulary "
        "(see descriptor 2).\n"
        "- Silence or answers far too short to judge an aspect pull that aspect down.\n"
        "- Comprehension/Interaction: did each answer address its question? Use the repeats used and the "
        "response latency as evidence.\n"
        "- Score the whole test (all questions together), one integer 1-5 per aspect.\n"
        "- evidence: 1-2 short sentences in Indonesian quoting or describing what you heard.\n"
        "- overall_comment: 2-3 sentences in Indonesian for the researcher (not for the student).\n"
        "Return only the JSON object.\n\nTHE TEST:"
    )}]
    for item in all_items(content):
        resp = by_item.get(item["id"])
        header = f"\n[{item['part_title']}] Question {item['id']}: \"{item['prompt']}\""
        if not resp:
            parts.append({"text": header + "\n(no answer recorded)"})
            continue
        meta = resp.get("transcript_meta") or {}
        info = (
            f"{header}\nRepeats used: {resp.get('repeats_used', 0)}. "
            f"Answer length: {meta.get('duration_s', resp.get('duration_seconds'))} s of {item.get('speak_seconds')} s allowed. "
            f"Words: {meta.get('words', '?')}, speech rate: {meta.get('speech_rate_wpm', '?')} wpm, "
            f"pauses >1s: {meta.get('pauses_over_1s', '?')}, longest pause: {meta.get('longest_pause_s', '?')} s, "
            f"response latency: {meta.get('response_latency_s', '?')} s, "
            f"Indonesian words: {meta.get('indonesian_words', '?')}.\n"
            f"Automatic transcript: \"{(resp.get('transcript') or '').strip() or '(nothing recognised)'}\"\n"
            "Audio:"
        )
        parts.append({"text": info})
        audio = audio_by_item.get(item["id"])
        if audio:
            mime = (resp.get("mime_type") or "audio/webm").split(";")[0]
            parts.append({"inlineData": {"mimeType": _GEMINI_MIME.get(mime, mime),
                                         "data": base64.b64encode(audio).decode()}})
    return parts


def _call_gemini(parts):
    key = os.getenv("GEMINI_API_KEY")
    if not key:
        raise RuntimeError("GEMINI_API_KEY is not set on the server.")
    body = {
        "contents": [{"role": "user", "parts": parts}],
        "generationConfig": {
            "temperature": 0, "responseMimeType": "application/json", "responseSchema": SCORE_SCHEMA,
            "maxOutputTokens": 8192,
        },
    }
    with _client(180) as client:
        resp = client.post(f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent",
                           headers={"x-goog-api-key": key, "Content-Type": "application/json"}, json=body)
    if resp.status_code == 429:
        raise _handle_429(resp)
    if resp.status_code >= 500:
        raise RuntimeError(f"Gemini server error {resp.status_code}")
    _check(resp, "Gemini scoring")
    data = resp.json()
    try:
        text = "".join(p.get("text", "") for p in data["candidates"][0]["content"]["parts"] if not p.get("thought"))
        result = json.loads(text)
    except Exception:
        raise RuntimeError(f"Gemini returned no usable score: {json.dumps(data)[:300]}")
    run = {}
    for a in ASPECTS:
        score = int((result.get(a) or {}).get("score"))
        if not 1 <= score <= 5:
            raise RuntimeError(f"Gemini gave {a}={score}, outside 1-5")
        run[a] = score
        run[f"{a}_evidence"] = str((result.get(a) or {}).get("evidence") or "")[:600]
    run["overall_comment"] = str(result.get("overall_comment") or "")[:1500]
    usage = data.get("usageMetadata") or {}
    run["tokens"] = usage.get("totalTokenCount")
    return run


def final_scores(runs):
    scores = {}
    for a in ASPECTS:
        vals = sorted(int(x[a]) for x in runs)
        scores[a] = vals[(len(vals) - 1) // 2]  # median (lower median for an even count)
    total = sum(scores[a] for a in ASPECTS)
    scores["total"] = total
    scores["nilai"] = round(total / 25 * 100, 2)
    return scores


def attempt_metrics(responses):
    metas = [r_.get("transcript_meta") or {} for r_ in responses]
    words = sum(m.get("words", 0) or 0 for m in metas)
    dur = sum(m.get("duration_s", 0) or 0 for m in metas)
    indo = sum(m.get("indonesian_words", 0) or 0 for m in metas)
    lat = [m["response_latency_s"] for m in metas if isinstance(m.get("response_latency_s"), (int, float))]
    return {
        "answers": len(responses),
        "total_speaking_s": round(dur, 1),
        "words": words,
        "speech_rate_wpm": round(words / (dur / 60), 1) if dur else 0,
        "pauses_over_1s": sum(m.get("pauses_over_1s", 0) or 0 for m in metas),
        "pauses_over_2s": sum(m.get("pauses_over_2s", 0) or 0 for m in metas),
        "longest_pause_s": max([m.get("longest_pause_s", 0) or 0 for m in metas] or [0]),
        "mean_response_latency_s": round(statistics.mean(lat), 2) if lat else None,
        "indonesian_words": indo,
        "indonesian_ratio": round(indo / words, 3) if words else 0,
        "repeats_used": sum(int(r_.get("repeats_used") or 0) for r_ in responses),
        "transcribers": sorted({m.get("model") for m in metas if m.get("model")}),
    }


def job_score(attempt_id):
    try:
        _score(attempt_id)
    finally:
        _unlock(f"score:{attempt_id}")


def _score(attempt_id):
    attempt = db_one("research_attempts", {"id": f"eq.{attempt_id}", "select": "*"})
    if not attempt or attempt["status"] not in ("submitted", "scoring"):
        return
    test = db_one("research_tests", {"id": f"eq.{attempt['test_id']}", "select": "content"})
    responses = db_select("research_responses", {"attempt_id": f"eq.{attempt_id}", "select": "*"})
    if any(x["transcribe_status"] == "pending" for x in responses):
        return  # the sweeper comes back when every answer is transcribed

    # Keep finished runs (also across days) as long as model and prompt are the same.
    runs = [x for x in (attempt.get("ai_runs") or [])
            if x.get("model") == GEMINI_MODEL and x.get("prompt_version") == PROMPT_VERSION]
    failures = int(r.get(f"research:fail:{attempt_id}") or 0)
    db_update("research_attempts", {"id": f"eq.{attempt_id}"},
              {"status": "scoring", "metrics": attempt_metrics(responses)})

    audio = {}
    if len(runs) < SCORING_RUNS:
        with _client(60) as client:
            for resp in responses:
                try:
                    audio[resp["item_id"]] = storage_download(resp["audio_path"], client)
                except Exception as e:
                    print(f"[research] could not download {resp['audio_path']}: {e}")
        parts = _scoring_parts(test.get("content"), responses, audio)

    while len(runs) < SCORING_RUNS:
        wait = _gemini_acquire()
        if wait > 0 and wait < 90:
            time.sleep(wait)
            continue
        if wait > 0:
            _wait_for_quota(attempt_id, runs, wait)
            return
        try:
            run = _call_gemini(parts)
        except _GeminiQuota as q:
            if not q.daily and q.wait < 90:
                time.sleep(q.wait)
                continue
            _wait_for_quota(attempt_id, runs, q.wait)
            return
        except Exception as e:
            failures += 1
            r.set(f"research:fail:{attempt_id}", failures, ex=7 * 86400)
            print(f"[research] scoring {attempt_id} failed ({failures}): {e}")
            if failures >= MAX_SCORE_FAILURES:
                db_update("research_attempts", {"id": f"eq.{attempt_id}"},
                          {"status": "error", "ai_runs": runs, "ai_error": str(e)[:400], "ai_next_try_at": None})
            else:
                db_update("research_attempts", {"id": f"eq.{attempt_id}"}, {
                    "status": "submitted", "ai_runs": runs, "ai_error": f"Retrying soon: {str(e)[:300]}",
                    "ai_next_try_at": (datetime.now(timezone.utc) + timedelta(minutes=5 * failures)).isoformat()})
            return
        run.update({"model": GEMINI_MODEL, "prompt_version": PROMPT_VERSION, "at": now_iso()})
        runs.append(run)
        db_update("research_attempts", {"id": f"eq.{attempt_id}"}, {"ai_runs": runs})

    scores = final_scores(runs)
    # Evidence/comment from the run closest to the final scores.
    best = min(runs, key=lambda x: sum(abs(x[a] - scores[a]) for a in ASPECTS))
    feedback = {a: best.get(f"{a}_evidence", "") for a in ASPECTS}
    feedback["overall_comment"] = best.get("overall_comment", "")
    db_update("research_attempts", {"id": f"eq.{attempt_id}"}, {
        "status": "scored", "ai_scores": scores, "ai_runs": runs, "ai_feedback": feedback,
        "ai_model": GEMINI_MODEL, "ai_prompt_version": PROMPT_VERSION, "ai_scored_at": now_iso(),
        "ai_error": None, "ai_next_try_at": None,
    })
    r.delete(f"research:fail:{attempt_id}")


def _wait_for_quota(attempt_id, runs, wait):
    until = datetime.now(timezone.utc) + timedelta(seconds=max(60, wait))
    local = until.astimezone(LOCAL_TZ).strftime("%d %b %H:%M")
    db_update("research_attempts", {"id": f"eq.{attempt_id}"}, {
        "status": "submitted", "ai_runs": runs,
        "ai_error": f"Waiting for today's free AI scoring quota (continues automatically after {local} WITA).",
        "ai_next_try_at": until.isoformat()})


# =============================================================================
# Sweeper: (re)start pending work. Called every couple of minutes by the API.
# =============================================================================

def sweep(queue_transcribe, queue_score):
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    started = {"transcribe": 0, "score": 0}
    with _client() as client:
        pending = db_select("research_responses", {
            "select": "id,research_attempts!inner(status)",
            "transcribe_status": "eq.pending",
            "or": f"(next_try_at.is.null,next_try_at.lte.{now})",
            "research_attempts.status": "in.(submitted,scoring)",
            "order": "created_at.asc",
        }, client)
        for row in pending[:200]:
            if _lock(f"tr:{row['id']}", 3600):
                queue_transcribe.enqueue("utils.research.job_transcribe", row["id"],
                                         job_timeout=900, result_ttl=600, failure_ttl=86400)
                started["transcribe"] += 1

        attempts = db_select("research_attempts", {
            "select": "id", "status": "in.(submitted,scoring)",
            "or": f"(ai_next_try_at.is.null,ai_next_try_at.lte.{now})",
            "order": "submitted_at.asc",
        }, client)
        if attempts:
            waiting = db_select("research_responses", {
                "select": "attempt_id", "transcribe_status": "eq.pending",
                "attempt_id": in_list(a["id"] for a in attempts)}, client)
            not_ready = {w["attempt_id"] for w in waiting}
            for a in attempts:
                if a["id"] in not_ready:
                    continue
                if _lock(f"score:{a['id']}", 3 * 3600):
                    queue_score.enqueue("utils.research.job_score", a["id"],
                                        job_timeout=3 * 3600, result_ttl=600, failure_ttl=86400)
                    started["score"] += 1
    return started


# =============================================================================
# Capacity: how many students per day the free limits can handle
# =============================================================================

def capacity(test_content: Optional[dict] = None):
    from utils.ratelimit import LIMITS, SAFETY
    # --- transcription (Groq Whisper, shared with practice modes) ---
    max_answer_s = sum(i.get("speak_seconds", 30) for i in all_items(test_content)) if test_content else 330
    n_items = len(all_items(test_content)) if test_content else 8
    observed = None
    try:
        rows = db_select("research_responses", {"select": "attempt_id,duration_seconds", "order": "created_at.desc",
                                                "limit": 400})
        per_attempt = {}
        for row in rows:
            d = float(row.get("duration_seconds") or 0)
            per_attempt[row["attempt_id"]] = per_attempt.get(row["attempt_id"], 0) + max(10, d)
        if len(per_attempt) >= 5:
            observed = statistics.mean(per_attempt.values())
    except Exception:
        pass
    # Groq bills at least 10 s per recording.
    est_audio_per_student = observed or max(n_items * 10, max_answer_s * 0.8)

    now = time.time()
    day_bucket = int(now // 86400)
    hour_bucket = int(now // 3600)
    whisper = []
    for model in ai.TRANSCRIBE_MODELS:
        lim = LIMITS.get(model, {})
        asd = int(lim.get("asd", 28800) * SAFETY)
        ash = int(lim.get("ash", 7200) * SAFETY)
        used_day = int(r.get(f"rl:{model}:asd:{day_bucket}") or 0)
        used_hour = int(r.get(f"rl:{model}:ash:{hour_bucket}") or 0)
        whisper.append({"model": model, "per_day_s": asd, "per_hour_s": ash,
                        "used_today_s": used_day, "used_this_hour_s": used_hour})
    w_day = sum(w["per_day_s"] for w in whisper)
    w_hour = sum(w["per_hour_s"] for w in whisper)
    w_left = sum(max(0, w["per_day_s"] - w["used_today_s"]) for w in whisper)
    groq_reset = datetime.fromtimestamp((day_bucket + 1) * 86400, timezone.utc)

    # --- AI scoring (Gemini) ---
    rpd = gemini_rpd()
    used = gemini_used_today()
    blocked = _gemini_blocked_until() > now
    g_day = rpd // SCORING_RUNS
    g_left = 0 if blocked else max(0, rpd - used) // SCORING_RUNS

    # --- backlog ---
    try:
        backlog_tr = len(db_select("research_responses", {"select": "id", "transcribe_status": "eq.pending"}))
        backlog_score = len(db_select("research_attempts", {"select": "id", "status": "in.(submitted,scoring)"}))
        storage_bytes = sum(int(x.get("size_bytes") or 0)
                            for x in db_select("research_responses", {"select": "size_bytes"}))
    except Exception:
        backlog_tr = backlog_score = storage_bytes = None

    transcribe_students_day = int(w_day // est_audio_per_student)
    recommended = max(1, min(transcribe_students_day, g_day))
    to_local = lambda dt: dt.astimezone(LOCAL_TZ).strftime("%H:%M")  # noqa: E731
    return {
        "recommended_students_per_day": recommended,
        "bottleneck": "AI scoring (Gemini)" if g_day <= transcribe_students_day else "transcription (Groq)",
        "transcription": {
            "models": whisper,
            "est_audio_seconds_per_student": round(est_audio_per_student),
            "estimate_from": "recorded tests" if observed else "test timings",
            "students_per_day": transcribe_students_day,
            "students_per_hour": int(w_hour // est_audio_per_student),
            "students_left_today": int(w_left // est_audio_per_student),
            "resets_at_local": to_local(groq_reset),
            "note": "Shared with practice modes (Read Aloud etc.) - heavy practice use lowers this.",
        },
        "scoring": {
            "model": GEMINI_MODEL,
            "requests_per_day": rpd,
            "requests_learned_from_google": bool(r.get(f"research:gemini:{GEMINI_MODEL}:rpd")),
            "used_today": used,
            "runs_per_test": SCORING_RUNS,
            "students_per_day": g_day,
            "students_left_today": g_left,
            "quota_used_up": blocked,
            "resets_at_local": to_local(next_gemini_reset()),
        },
        "backlog": {
            "answers_waiting_transcription": backlog_tr,
            "tests_waiting_scoring": backlog_score,
            "days_to_clear": (math.ceil(backlog_score / g_day) if backlog_score and g_day else 0)
            if backlog_score is not None else None,
        },
        "storage": {
            "used_mb": round(storage_bytes / 1048576, 1) if storage_bytes is not None else None,
            "limit_mb": STORAGE_LIMIT_MB,
            "est_mb_per_student_test": round(est_audio_per_student * 4200 / 1048576, 2),
        },
        "local_tz": str(LOCAL_TZ),
    }


# =============================================================================
# Export (Excel) and recordings (ZIP)
# =============================================================================

def _pearson(xs, ys):
    n = len(xs)
    if n < 3:
        return None
    mx, my = statistics.mean(xs), statistics.mean(ys)
    sx = math.sqrt(sum((x - mx) ** 2 for x in xs))
    sy = math.sqrt(sum((y - my) ** 2 for y in ys))
    if sx == 0 or sy == 0:
        return None
    return round(sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / (sx * sy), 3)


def build_export(study, tests, attempts, students, classes, human, scorers, responses):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4F46E5")
    stud = {s["id"]: s for s in students}
    cls = {c["id"]: c.get("class_name", "") for c in classes}
    test_kind = {t["id"]: t["kind"] for t in tests}
    scorer_name = {s["user_id"]: s["full_name"] for s in scorers}
    kind_label = {"pre": "Pretest", "post": "Posttest"}
    short = ["Pron.", "Flu.", "Voc.", "Gram.", "Comp."]

    def sheet(wb, title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            ws.append(row)
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(9, min(40, len(str(h)) + 4)))
        ws.freeze_panes = "A2"
        return ws

    wb = Workbook()
    wb.remove(wb.active)
    by_student = {}
    for a in attempts:
        by_student.setdefault(a["student_id"], {})[test_kind.get(a["test_id"])] = a

    def sort_key(sid):
        s = stud.get(sid, {})
        return (cls.get(s.get("class_id"), ""), s.get("full_name", ""))

    # 1. Pre vs post summary (AI = main score)
    rows = []
    for no, sid in enumerate(sorted(by_student, key=sort_key), 1):
        s = stud.get(sid, {})
        row = [no, s.get("full_name", "(removed student)"), cls.get(s.get("class_id"), "")]
        nilai = {}
        for kind in ("pre", "post"):
            a = by_student[sid].get(kind)
            sc = (a or {}).get("ai_scores") or {}
            row += [sc.get(x) for x in ASPECTS] + [sc.get("total"), sc.get("nilai")]
            nilai[kind] = sc.get("nilai")
        gain = round(nilai["post"] - nilai["pre"], 2) if nilai["pre"] is not None and nilai["post"] is not None else None
        row.append(gain)
        rows.append(row)
    header = ["No.", "Nama Siswa", "Kelas"]
    for kind in ("pre", "post"):
        header += [f"{kind_label[kind]} {x}" for x in short] + [f"{kind_label[kind]} Total", f"{kind_label[kind]} Nilai"]
    header.append("Gain (Nilai)")
    sheet(wb, "Rekap Pre-Post (AI)", header, rows, {2: 28})

    # 2/3. Lembar Rekap Skor style, one sheet per test
    for t in sorted(tests, key=lambda x: x["kind"] != "pre"):  # pretest first
        rows = []
        tas = [a for a in attempts if a["test_id"] == t["id"]]
        for no, a in enumerate(sorted(tas, key=lambda a: sort_key(a["student_id"])), 1):
            s = stud.get(a["student_id"], {})
            sc = a.get("ai_scores") or {}
            rows.append([no, s.get("full_name", ""), cls.get(s.get("class_id"), "")]
                        + [sc.get(x) for x in ASPECTS] + [sc.get("total"), sc.get("nilai"), a["status"]])
        sheet(wb, f"Rekap {kind_label[t['kind']]}",
              ["No.", "Nama Siswa", "Kelas"] + short + ["Total", "Nilai", "Status"], rows, {2: 28})

    # 4. AI vs human
    att = {a["id"]: a for a in attempts}
    rows, pairs = [], {a: ([], []) for a in ASPECTS + ["total"]}
    for h in human:
        a = att.get(h["attempt_id"])
        if not a:
            continue
        s = stud.get(a["student_id"], {})
        ai_sc = a.get("ai_scores") or {}
        h_total = sum(int(h[x]) for x in ASPECTS)
        rows.append([a["anon_code"], s.get("full_name", ""), kind_label.get(test_kind.get(a["test_id"]), ""),
                     scorer_name.get(h["scorer_id"], h["scorer_id"])]
                    + [ai_sc.get(x) for x in ASPECTS] + [ai_sc.get("total")]
                    + [h[x] for x in ASPECTS] + [h_total, round(h_total / 25 * 100, 2), h.get("notes") or ""])
        if ai_sc:
            for x in ASPECTS:
                pairs[x][0].append(int(ai_sc[x]))
                pairs[x][1].append(int(h[x]))
            pairs["total"][0].append(int(ai_sc["total"]))
            pairs["total"][1].append(h_total)
    sheet(wb, "AI vs Penilai Manusia",
          ["Kode", "Nama Siswa", "Tes", "Penilai"] + [f"AI {x}" for x in short] + ["AI Total"]
          + [f"Manusia {x}" for x in short] + ["Manusia Total", "Manusia Nilai", "Catatan"], rows, {2: 26, 18: 40})

    # 5. Agreement
    rows = []
    for x in ASPECTS + ["total"]:
        xs, ys = pairs[x]
        n = len(xs)
        if n == 0:
            rows.append([ASPECT_LABELS.get(x, "Total (5-25)"), 0, None, None, None, None])
            continue
        exact = sum(1 for a_, b_ in zip(xs, ys) if a_ == b_) / n
        within = sum(1 for a_, b_ in zip(xs, ys) if abs(a_ - b_) <= (1 if x != "total" else 2)) / n
        mean_diff = statistics.mean(a_ - b_ for a_, b_ in zip(xs, ys))
        rows.append([ASPECT_LABELS.get(x, "Total (5-25)"), n, _pearson(xs, ys), round(exact * 100, 1),
                     round(within * 100, 1), round(mean_diff, 2)])
    ws = sheet(wb, "Kesepakatan AI-Manusia",
               ["Aspek", "n (pasangan)", "Pearson r", "Sama persis (%)", "Selisih <=1 (Total <=2) (%)",
                "Rata-rata selisih (AI - manusia)"], rows, {1: 30})
    ws.append([])
    ws.append(["Pearson r and agreement between the AI (main scorer) and human scorers, per scored pair."])

    # 6. Audio measures
    rows = []
    for a in sorted(attempts, key=lambda a: (test_kind.get(a["test_id"], ""), sort_key(a["student_id"]))):
        m = a.get("metrics") or {}
        s = stud.get(a["student_id"], {})
        rows.append([s.get("full_name", ""), cls.get(s.get("class_id"), ""), kind_label.get(test_kind.get(a["test_id"]), ""),
                     a["anon_code"], m.get("total_speaking_s"), m.get("words"), m.get("speech_rate_wpm"),
                     m.get("pauses_over_1s"), m.get("pauses_over_2s"), m.get("longest_pause_s"),
                     m.get("mean_response_latency_s"), m.get("indonesian_words"), m.get("indonesian_ratio"),
                     m.get("repeats_used"), ", ".join(m.get("transcribers") or [])])
    sheet(wb, "Ukuran Audio",
          ["Nama Siswa", "Kelas", "Tes", "Kode", "Durasi bicara (s)", "Jumlah kata", "Kecepatan (kata/menit)",
           "Jeda >1s", "Jeda >2s", "Jeda terpanjang (s)", "Rata-rata waktu mulai menjawab (s)",
           "Kata bahasa Indonesia", "Rasio bahasa Indonesia", "Pengulangan soal", "Transkripsi oleh"], rows, {1: 26})

    # 7. AI details (reproducibility)
    rows = []
    for a in sorted(attempts, key=lambda a: (test_kind.get(a["test_id"], ""), sort_key(a["student_id"]))):
        s = stud.get(a["student_id"], {})
        fb = a.get("ai_feedback") or {}
        runs = a.get("ai_runs") or []
        rows.append([s.get("full_name", ""), kind_label.get(test_kind.get(a["test_id"]), ""), a["anon_code"], a["status"],
                     a.get("ai_model"), a.get("ai_prompt_version"), len(runs),
                     " | ".join("/".join(str(x.get(k)) for k in ASPECTS) for x in runs),
                     a.get("ai_scored_at"), fb.get("overall_comment", ""),
                     *[fb.get(x, "") for x in ASPECTS], a.get("ai_error") or ""])
    sheet(wb, "Detail AI",
          ["Nama Siswa", "Tes", "Kode", "Status", "Model", "Versi prompt", "Jumlah run", "Skor tiap run (P/F/V/G/C)",
           "Dinilai pada (UTC)", "Komentar AI"] + [f"Bukti {x}" for x in short] + ["Masalah"], rows, {1: 26, 10: 50})

    # 8. Consent
    rows = []
    for a in sorted(attempts, key=lambda a: (test_kind.get(a["test_id"], ""), sort_key(a["student_id"]))):
        s = stud.get(a["student_id"], {})
        rows.append([s.get("full_name", ""), cls.get(s.get("class_id"), ""), kind_label.get(test_kind.get(a["test_id"]), ""),
                     "Setuju" if a.get("consent_given") else "Tidak setuju", a.get("consent_at"),
                     a.get("started_at"), a.get("submitted_at")])
    sheet(wb, "Persetujuan", ["Nama Siswa", "Kelas", "Tes", "Persetujuan", "Waktu persetujuan (UTC)",
                              "Mulai (UTC)", "Selesai (UTC)"], rows, {1: 26})

    # 9. Transcripts
    rows = []
    for x in sorted(responses, key=lambda x: (x.get("attempt_id"), x.get("item_id"))):
        a = att.get(x["attempt_id"])
        if not a:
            continue
        s = stud.get(a["student_id"], {})
        meta = x.get("transcript_meta") or {}
        rows.append([s.get("full_name", ""), kind_label.get(test_kind.get(a["test_id"]), ""), a["anon_code"],
                     x["item_id"], x.get("duration_seconds"), x.get("repeats_used"), meta.get("language"),
                     x.get("transcript") or ""])
    sheet(wb, "Transkrip", ["Nama Siswa", "Tes", "Kode", "Soal", "Durasi (s)", "Pengulangan", "Bahasa terdeteksi",
                            "Transkrip"], rows, {1: 26, 8: 80})

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _safe(name):
    return re.sub(r"[^\w\- .()]+", "_", str(name or "")).strip()[:60] or "unnamed"


def build_zip(path_out, tests, attempts, students, classes, responses):
    stud = {s["id"]: s for s in students}
    cls = {c["id"]: c.get("class_name", "") for c in classes}
    kind = {t["id"]: ("Pretest" if t["kind"] == "pre" else "Posttest") for t in tests}
    att = {a["id"]: a for a in attempts}
    count = 0
    with zipfile.ZipFile(path_out, "w", compression=zipfile.ZIP_STORED) as z, _client(120) as client:
        lines = ["test,class,student,code,item,file,duration_s,repeats,transcript"]
        for x in responses:
            a = att.get(x["attempt_id"])
            if not a:
                continue
            s = stud.get(a["student_id"], {})
            ext = os.path.splitext(x["audio_path"])[1] or ".webm"
            folder = f"{kind.get(a['test_id'], 'Test')}/{_safe(cls.get(s.get('class_id')))}/{_safe(s.get('full_name'))} ({a['anon_code']})"
            name = f"{folder}/{x['item_id']}{ext}"
            try:
                z.writestr(name, storage_download(x["audio_path"], client))
                count += 1
            except Exception as e:
                print(f"[research] zip: could not add {x['audio_path']}: {e}")
                continue
            t = (x.get("transcript") or "").replace('"', '""')
            lines.append(f"{kind.get(a['test_id'])},\"{cls.get(s.get('class_id'), '')}\",\"{s.get('full_name', '')}\","
                         f"{a['anon_code']},{x['item_id']},\"{name}\",{x.get('duration_seconds') or ''},"
                         f"{x.get('repeats_used') or 0},\"{t}\"")
        z.writestr("index.csv", "﻿" + "\n".join(lines))
    return count
