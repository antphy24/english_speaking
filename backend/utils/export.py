"""
Teacher score export -> Excel (.xlsx).

Built on the server so it always contains ALL matching records (no 1000-row or
browser limits), with numeric score columns that Excel can sort/average, and a
per-student summary sheet for grading.
"""
import io
import os
import re
from collections import defaultdict
from datetime import datetime, timedelta, timezone

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PAGE = 1000
MODE_LABELS = {
    "read_aloud": "Read Aloud",
    "qa": "Q&A",
    "conversation": "Conversation",
    "debate": "Debate",
    "mattering": "Mattering",
}
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# --- Supabase access (service role; the caller has already verified the teacher) ---

def _client():
    force_ipv4 = os.getenv("FORCE_IPV4", "true").lower() == "true"
    transport = httpx.HTTPTransport(local_address="0.0.0.0", retries=2) if force_ipv4 else None
    return httpx.Client(timeout=60, transport=transport)


def _headers():
    key = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
    return {"apikey": key, "Authorization": f"Bearer {key}"}


def _get_all(client, path, params):
    """GET every page of a PostgREST query."""
    url = f"{os.getenv('SUPABASE_URL')}/rest/v1/{path}"
    rows, offset = [], 0
    while True:
        resp = client.get(url, headers=_headers(), params={**params, "limit": PAGE, "offset": offset})
        if resp.status_code != 200:
            raise RuntimeError(f"Supabase query on {path} failed ({resp.status_code}): {resp.text[:300]}")
        batch = resp.json()
        rows.extend(batch)
        if len(batch) < PAGE:
            return rows
        offset += PAGE


def fetch_export_data(teacher_id, class_id=None, mode=None, days=None):
    with _client() as client:
        classes = _get_all(client, "classes", {
            "select": "id,class_name", "teacher_id": f"eq.{teacher_id}", "order": "class_name.asc"})
        if class_id and class_id != "all":
            classes = [c for c in classes if c["id"] == class_id]
            if not classes:
                raise PermissionError("You do not own this class.")
        if not classes:
            return classes, [], []
        class_ids = ",".join(c["id"] for c in classes)

        students = _get_all(client, "students", {
            "select": "id,full_name,school_id,class_id", "class_id": f"in.({class_ids})",
            "order": "full_name.asc"})

        base = {
            "students.class_id": f"in.({class_ids})",
            "order": "created_at.desc",
        }
        if mode and mode != "all":
            base["mode"] = f"eq.{mode}"
        if days:
            since = datetime.now(timezone.utc) - timedelta(days=int(days))
            base["created_at"] = f"gte.{since.isoformat()}"
        cols = "id,student_id,mode,score,feedback,created_at,duration_seconds"
        embed = "students!inner(class_id)"
        try:
            assessments = _get_all(client, "assessments", {**base, "select": f"{cols},transcript,{embed}"})
        except RuntimeError as e:
            if "transcript" not in str(e):
                raise
            # The optional 'transcript' column migration has not been run yet.
            assessments = _get_all(client, "assessments", {**base, "select": f"{cols},{embed}"})
    return classes, students, assessments


def fetch_material_dates(classes):
    """
    When each of the teacher's materials was created: {(mode, material name): ISO time}.
    Includes materials shared with all classes. A material given to several classes
    counts from its earliest copy. Used to order the score-matrix columns.
    """
    class_ids = ",".join(c["id"] for c in classes)
    scope = f"(class_id.in.({class_ids}),class_id.is.null)" if class_ids else "(class_id.is.null)"
    with _client() as client:
        rows = _get_all(client, "custom_materials", {
            "select": "mode,title,content,created_at", "or": scope, "order": "created_at.asc"})
    created = {}
    for row in rows:
        when = row.get("created_at")
        if not when:
            continue
        content = (row.get("content") or "").strip()
        # Scores are filed under the title; debate files them under the motion, and a
        # mattering issue without a title under its first line (the motion).
        names = {(row.get("title") or "").strip(), content, content.split("\n")[0].strip()}
        for name in names:
            key = (row.get("mode"), name)
            if name and (key not in created or when < created[key]):
                created[key] = when
    return created


# The app's built-in materials, in the order students see them (easy -> hard).
# Keep in step with the lists at the top of the frontend Mode*.jsx files.
BUILTIN_MATERIALS = {
    "read_aloud": ["Vocal Warmup (Easy)", "Climate Change (Medium)", "Digital Technology (Hard)"],
    "qa": ["Problem Solving", "Memorable Journey", "Healthy Hobby"],
    "conversation": ["General Chat", "Job Interview Practice", "Travel & Tourism"],
    "debate": [
        "This House would ban the use of AI in educational assessments.",
        "This House believes that developing nations should prioritize economic growth over environmental protection.",
        "This House would implement a 4-day work week.",
        "This House regrets the rise of cancel culture.",
    ],
    "mattering": ["Phones in school", "Homework", "Social media age limit", "Single-use plastic"],
}


def material_order_key(mode, name, created, first_used):
    """
    Column order in the score matrix:
      1. built-in materials, in their own order (easy, medium, hard);
      2. the teacher's materials, oldest created first;
      3. anything else (deleted / renamed materials, old untitled attempts), by first use.
    """
    builtin = BUILTIN_MATERIALS.get(mode, [])
    if name in builtin:
        return (0, builtin.index(name), name)
    if created:
        return (1, _when(created), name)
    return (2, _when(first_used), name)


def _when(iso):
    try:
        return datetime.fromisoformat(str(iso).replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return datetime.max.replace(tzinfo=timezone.utc)


# --- helpers ------------------------------------------------------------------

def material_of(feedback):
    fb = feedback if isinstance(feedback, dict) else {}
    return fb.get("material_title") or fb.get("motion") or "Default Material"


def _num(value, digits=None):
    try:
        v = float(value)
    except (TypeError, ValueError):
        return None
    return round(v, digits) if digits is not None else v


def _text(value):
    if value is None:
        return None
    if isinstance(value, list):
        value = ", ".join(str(v) for v in value)
    value = _ILLEGAL.sub("", str(value))
    return value[:32000]  # Excel cell limit is 32767 characters


def _local(iso, tz_offset_minutes):
    try:
        dt = datetime.fromisoformat(iso.replace("Z", "+00:00"))
    except (AttributeError, ValueError):
        return None
    # JS getTimezoneOffset(): minutes to ADD to local time to get UTC (UTC+8 -> -480)
    return (dt - timedelta(minutes=tz_offset_minutes or 0)).replace(tzinfo=None)


def _style_sheet(ws, widths, header_row=1):
    header_fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    if ws.max_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"


def _append(ws, values):
    ws.append(values)
    for cell in ws[ws.max_row]:
        # Never let student/AI text be interpreted as an Excel formula.
        if isinstance(cell.value, str) and cell.value[:1] in ("=", "+", "-", "@"):
            cell.data_type = "s"


# --- workbook -------------------------------------------------------------------

def build_workbook(classes, students, assessments, *, mode=None, material=None, search=None,
                   days=None, tz_offset_minutes=0, teacher_name=""):
    class_name = {c["id"]: c["class_name"] for c in classes}
    student_by_id = {s["id"]: s for s in students}

    q = (search or "").strip().lower()

    def student_matches(s):
        return not q or q in (s.get("full_name") or "").lower() or q in (s.get("school_id") or "").lower()

    rows = []
    for a in assessments:
        s = student_by_id.get(a["student_id"])
        if not s or not student_matches(s):
            continue
        if material and material != "all" and material_of(a.get("feedback")) != material:
            continue
        rows.append((a, s))

    wb = Workbook()

    # Sheet 1: one row per student (for grading)
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Student", "School ID", "Class", "Mode", "Attempts", "Best score", "Average score",
               "Latest score", "First score", "Change (latest - first)", "Latest attempt"])
    per_student = defaultdict(list)
    for a, s in rows:
        per_student[(s["id"], a["mode"])].append(a)
    summary_students = sorted((s for s in students if student_matches(s)),
                              key=lambda s: (class_name.get(s["class_id"], ""), (s.get("full_name") or "").lower()))
    modes = [mode] if mode and mode != "all" else list(MODE_LABELS)
    for s in summary_students:
        student_rows = 0
        for m in modes:
            attempts = sorted(per_student.get((s["id"], m), []), key=lambda a: a["created_at"])
            scores = [_num(a.get("score")) for a in attempts if _num(a.get("score")) is not None]
            if not scores:
                continue
            student_rows += 1
            _append(ws, [
                _text(s.get("full_name")), _text(s.get("school_id")), class_name.get(s["class_id"]),
                MODE_LABELS.get(m, m), len(attempts), max(scores), round(sum(scores) / len(scores), 1),
                scores[-1], scores[0], scores[-1] - scores[0],
                _local(attempts[-1]["created_at"], tz_offset_minutes),
            ])
        if student_rows == 0:  # show who has not done it yet
            _append(ws, [_text(s.get("full_name")), _text(s.get("school_id")), class_name.get(s["class_id"]),
                         MODE_LABELS.get(mode, "-") if mode and mode != "all" else "-",
                         0, None, None, None, None, None, "Not attempted"])
    for row in ws.iter_rows(min_row=2, min_col=11, max_col=11):
        row[0].number_format = "yyyy-mm-dd hh:mm"
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        row[0].number_format = "0.0"
    _style_sheet(ws, [28, 14, 16, 13, 10, 11, 13, 12, 11, 12, 18])

    # Sheet 2: every attempt with all details
    ws2 = wb.create_sheet("All attempts")
    ws2.append([
        "Date", "Student", "School ID", "Class", "Mode", "Material", "Score (0-100)",
        "Accuracy %", "Word error rate %", "Skipped words", "Unclear words",
        "Fluency", "Lexical resource", "Grammar", "Pronunciation", "Interactive communication",
        "Matter (1-10)", "Manner (1-10)", "Method (1-10)",
        "Feedback", "Feedback source", "Transcribed by", "Session time (min)", "What the student said (transcript)",
    ])
    for a, s in rows:
        fb = a.get("feedback") if isinstance(a.get("feedback"), dict) else {}
        wer = _num(fb.get("word_error_rate"))
        duration = _num(a.get("duration_seconds"))
        feedback_text = fb.get("feedback") or fb.get("overall_feedback")
        if a["mode"] == "debate":
            parts = [fb.get("overall_feedback"),
                     fb.get("matter_feedback") and f"Matter: {fb['matter_feedback']}",
                     fb.get("manner_feedback") and f"Manner: {fb['manner_feedback']}",
                     fb.get("method_feedback") and f"Method: {fb['method_feedback']}"]
            feedback_text = "\n".join(p for p in parts if p)
        if a["mode"] == "mattering":
            parts = [fb.get("band") and f"Band: {fb['band']}",
                     fb.get("overall_feedback"),
                     fb.get("argument_feedback") and f"Argument: {fb['argument_feedback']}",
                     fb.get("rebuttal_feedback") and f"Rebuttal: {fb['rebuttal_feedback']}",
                     fb.get("knowledge_gaps") and f"Read up on: {fb['knowledge_gaps']}"]
            feedback_text = "\n".join(p for p in parts if p)
        _append(ws2, [
            _local(a["created_at"], tz_offset_minutes),
            _text(s.get("full_name")), _text(s.get("school_id")), class_name.get(s["class_id"]),
            MODE_LABELS.get(a["mode"], a["mode"]), _text(material_of(fb)), _num(a.get("score")),
            _num(fb.get("accuracy_score")), round(wer * 100, 1) if wer is not None else None,
            _text(fb.get("skipped_words")), _text(fb.get("mispronounced_words")),
            _num(fb.get("fluency", fb.get("fluency_and_coherence"))), _num(fb.get("lexical_resource")),
            _num(fb.get("grammatical_range")), _num(fb.get("pronunciation")),
            _num(fb.get("interactive_communication")),
            _num(fb.get("matter_score")), _num(fb.get("manner_score")), _num(fb.get("method_score")),
            _text(feedback_text), {"ai": "AI", "standard": "Standard"}.get(fb.get("feedback_source"), "AI"),
            "Backup (server)" if str(fb.get("transcriber", "")).startswith("server:") else "Groq",
            round(duration / 60, 1) if duration else None,
            _text(a.get("transcript") or fb.get("transcript")),
        ])
    for row in ws2.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = "yyyy-mm-dd hh:mm"
    _style_sheet(ws2, [17, 26, 13, 14, 13, 28, 12, 11, 12, 28, 28, 9, 10, 9, 13, 13, 10, 10, 10, 70, 11, 13, 12, 70])

    # Sheet 3: what this file contains
    ws3 = wb.create_sheet("About")
    period = {None: "All time", "7": "Last 7 days", "30": "Last 30 days"}.get(str(days) if days else None,
                                                                          f"Last {days} days")
    info = [
        ("Exported", datetime.utcnow() - timedelta(minutes=tz_offset_minutes or 0)),
        ("Teacher", teacher_name or ""),
        ("Classes", ", ".join(c["class_name"] for c in classes) or "-"),
        ("Mode", MODE_LABELS.get(mode, "All modes") if mode and mode != "all" else "All modes"),
        ("Material", material if material and material != "all" else "All materials"),
        ("Period", period),
        ("Search", search or "-"),
        ("Attempts exported", len(rows)),
        ("Students in summary", len(summary_students)),
        ("", ""),
        ("Scores", "All scores are out of 100. Read Aloud = % of words read correctly. Q&A and Conversation = "
                   "average of the sub-scores. Debate = Matter 40% + Manner 40% + Method 20%."),
        ("Mattering scores", "Mattering uses the debate speaker scale instead: every score is between 69 and 81, "
                             "and 75 means an average speech. Do not average it with the other modes."),
        ("Feedback source", "AI = written by the AI; Standard = automatic paragraph used when many students "
                            "were waiting (Read Aloud only). Scores are identical either way."),
    ]
    for key, value in info:
        _append(ws3, [key, value])
    ws3["B1"].number_format = "yyyy-mm-dd hh:mm"
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 110
    for cell in ws3["A"]:
        cell.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows)


# --- Score matrix: students x materials, highest score, colour-coded ---------------

GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")


def build_matrix_workbook(classes, students, assessments, *, mode=None, material=None, search=None,
                          days=None, tz_offset_minutes=0, threshold=75, material_created=None):
    """
    One sheet per mode: a row per student, a column per material, each cell the
    student's HIGHEST score for that material. Cells are green at/above the pass
    mark and red below it; the pass mark is a cell at the top, so the teacher can
    change it in Excel and the colours update immediately.
    """
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Border, Side

    class_name = {c["id"]: c["class_name"] for c in classes}
    student_by_id = {s["id"]: s for s in students}
    q = (search or "").strip().lower()

    def student_matches(s):
        return not q or q in (s.get("full_name") or "").lower() or q in (s.get("school_id") or "").lower()

    # (mode) -> material -> student_id -> best score ; and first-use time per material
    best = defaultdict(lambda: defaultdict(dict))
    first_used = defaultdict(dict)
    for a in assessments:
        s = student_by_id.get(a["student_id"])
        score = _num(a.get("score"))
        if not s or score is None or not student_matches(s):
            continue
        mat = material_of(a.get("feedback"))
        if material and material != "all" and mat != material:
            continue
        m = a["mode"]
        cur = best[m][mat].get(s["id"])
        best[m][mat][s["id"]] = score if cur is None else max(cur, score)
        t = a.get("created_at") or ""
        if mat not in first_used[m] or t < first_used[m][mat]:
            first_used[m][mat] = t

    modes = [mode] if mode and mode != "all" else [m for m in MODE_LABELS if best.get(m)]
    if not modes:
        modes = [mode] if mode and mode != "all" else ["read_aloud"]

    rows_students = sorted((s for s in students if student_matches(s)),
                           key=lambda s: ((s.get("school_id") or "~"), (s.get("full_name") or "").lower()))

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1F2937")

    for m in modes:
        ws = wb.create_sheet(MODE_LABELS.get(m, m)[:31])
        # Built-in materials first, then the teacher's by creation date, then the rest.
        materials = sorted(best[m].keys(), key=lambda mat: material_order_key(
            m, mat, (material_created or {}).get((m, mat)), first_used[m].get(mat)))

        ws["A1"] = f"{MODE_LABELS.get(m, m)} - highest score per material"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = "Pass mark"
        ws["A2"].font = Font(bold=True)
        ws["B2"] = int(threshold)
        ws["B2"].fill = INPUT_FILL
        ws["B2"].font = Font(bold=True, size=12)
        ws["B2"].border = Border(left=Side(style="medium"), right=Side(style="medium"),
                                 top=Side(style="medium"), bottom=Side(style="medium"))
        ws["C2"] = "<- change this number; the colours update automatically. Green = at or above, red = below."
        ws["C2"].font = Font(italic=True, color="666666")

        header_row = 4
        first_mat_col = 4
        headers = ["School ID", "Student Name", "Class Name"] + materials + ["Materials done", "Average of best scores"]
        for col, value in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=_text(value))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center" if col >= first_mat_col else "left",
                                       vertical="center", wrap_text=True)
            cell.border = border
        ws.row_dimensions[header_row].height = 32

        last_mat_col = first_mat_col + len(materials) - 1
        done_col = first_mat_col + len(materials)
        avg_col = done_col + 1
        r = header_row
        for s in rows_students:
            r += 1
            scores = [best[m][mat].get(s["id"]) for mat in materials]
            present = [x for x in scores if x is not None]
            values = [_text(s.get("school_id")), _text(s.get("full_name")), class_name.get(s["class_id"])]
            values += [int(round(x)) if x is not None else None for x in scores]
            values += [len(present), round(sum(present) / len(present), 1) if present else None]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=r, column=col, value=value)
                cell.border = border
                if col >= first_mat_col:
                    cell.alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=avg_col).number_format = "0.0"
        last_row = r

        # Colour rules read the pass mark from $B$2 (blank cells stay uncoloured)
        if materials and last_row > header_row:
            first_letter = get_column_letter(first_mat_col)
            rng = f"{first_letter}{header_row + 1}:{get_column_letter(last_mat_col)}{last_row}"
            ref = f"{first_letter}{header_row + 1}"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f"AND(ISNUMBER({ref}),{ref}>=$B$2)"], fill=GREEN_FILL, stopIfTrue=True))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f"AND(ISNUMBER({ref}),{ref}<$B$2)"], fill=RED_FILL, stopIfTrue=True))
            # the average column follows the same pass mark
            avg_letter = get_column_letter(avg_col)
            avg_rng = f"{avg_letter}{header_row + 1}:{avg_letter}{last_row}"
            avg_ref = f"{avg_letter}{header_row + 1}"
            ws.conditional_formatting.add(avg_rng, FormulaRule(
                formula=[f"AND(ISNUMBER({avg_ref}),{avg_ref}>=$B$2)"], fill=GREEN_FILL, stopIfTrue=True))
            ws.conditional_formatting.add(avg_rng, FormulaRule(
                formula=[f"AND(ISNUMBER({avg_ref}),{avg_ref}<$B$2)"], fill=RED_FILL, stopIfTrue=True))

        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 14
        for col in range(first_mat_col, last_mat_col + 1):
            ws.column_dimensions[get_column_letter(col)].width = 13
        ws.column_dimensions[get_column_letter(done_col)].width = 11
        ws.column_dimensions[get_column_letter(avg_col)].width = 12
        ws.freeze_panes = ws.cell(row=header_row + 1, column=first_mat_col)
        if last_row > header_row:
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(avg_col)}{last_row}"
        # Printing: landscape, all columns on one page width, header row repeated
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{header_row}:{header_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), sum(len(v) for m in modes for v in best[m].values())
